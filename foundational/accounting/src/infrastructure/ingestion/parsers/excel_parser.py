"""
Excel file parser implementation.

Parses XLSX files using openpyxl library. Supports reading from
specific worksheets and handling various Excel data types.

Note: Requires openpyxl package: pip install openpyxl
"""

from datetime import datetime
from decimal import Decimal
from typing import Any, BinaryIO, TextIO

from src.infrastructure.ingestion.parsers.base import (
    BaseFileParser,
    FileFormat,
    FileParseError,
    ParsedRow,
    ParseResult,
)


class ExcelParser(BaseFileParser):
    """
    Parser for Excel XLSX files.

    Features:
    - Reads from specific worksheets or active sheet
    - Handles Excel data types (dates, numbers, booleans)
    - Skips empty rows automatically
    - Supports large files with read-only mode

    Requires: openpyxl package
    """

    def __init__(self, worksheet_name: str | None = None) -> None:
        """
        Initialize Excel parser.

        Args:
            worksheet_name: Name of worksheet to parse (None for active sheet)
        """
        self.worksheet_name = worksheet_name

    @property
    def supported_format(self) -> FileFormat:
        return FileFormat.EXCEL

    @property
    def supported_extensions(self) -> list[str]:
        return [".xlsx", ".xlsm"]

    def _get_file_mode(self) -> str:
        return "rb"  # Excel files are binary

    def _parse_stream(
        self,
        file: BinaryIO | TextIO,
        result: ParseResult,
        encoding: str,  # Ignored for Excel - binary format
        has_header: bool,
        skip_rows: int,
        max_rows: int | None,
        column_mapping: dict[str, str] | None,
    ) -> None:
        """Parse Excel file stream."""
        try:
            import openpyxl
        except ImportError:
            result.errors.append(
                FileParseError(
                    row_number=0,
                    column=None,
                    message="openpyxl package required for Excel parsing. "
                    "Install with: pip install openpyxl",
                    error_code="MISSING_DEPENDENCY",
                )
            )
            return

        try:
            # Use read_only mode for better memory efficiency
            workbook = openpyxl.load_workbook(
                file,
                read_only=True,
                data_only=True,  # Get values, not formulas
            )
        except Exception as e:
            result.errors.append(
                FileParseError(
                    row_number=0,
                    column=None,
                    message=f"Failed to open Excel file: {str(e)}",
                    error_code="EXCEL_OPEN_ERROR",
                )
            )
            return

        try:
            # Select worksheet
            if self.worksheet_name:
                if self.worksheet_name not in workbook.sheetnames:
                    result.errors.append(
                        FileParseError(
                            row_number=0,
                            column=None,
                            message=f"Worksheet '{self.worksheet_name}' not found. "
                            f"Available: {workbook.sheetnames}",
                            error_code="WORKSHEET_NOT_FOUND",
                        )
                    )
                    return
                worksheet = workbook[self.worksheet_name]
            else:
                worksheet = workbook.active

            if worksheet is None:
                result.errors.append(
                    FileParseError(
                        row_number=0,
                        column=None,
                        message="No active worksheet found",
                        error_code="NO_WORKSHEET",
                    )
                )
                return

            result.metadata["worksheet_name"] = worksheet.title
            result.metadata["available_worksheets"] = workbook.sheetnames

            # Parse rows
            headers: list[str] = []
            row_iterator = worksheet.iter_rows()
            current_row = 0

            # Handle header row
            if has_header:
                try:
                    header_row = next(row_iterator)
                    current_row = 1
                    headers = [
                        self._cell_to_string(cell) or f"column_{i}"
                        for i, cell in enumerate(header_row)
                    ]
                    # Remove trailing empty headers
                    while headers and not headers[-1].strip():
                        headers.pop()
                    result.headers = headers
                except StopIteration:
                    result.errors.append(
                        FileParseError(
                            row_number=0,
                            column=None,
                            message="Worksheet is empty",
                            error_code="EMPTY_WORKSHEET",
                        )
                    )
                    return

            # Skip additional rows
            for _ in range(skip_rows):
                try:
                    next(row_iterator)
                    current_row += 1
                except StopIteration:
                    break

            # Parse data rows
            rows_parsed = 0

            for row in row_iterator:
                current_row += 1

                if max_rows is not None and rows_parsed >= max_rows:
                    break

                # Skip empty rows
                if all(cell.value is None for cell in row):
                    continue

                parsed_row = self._parse_row(
                    row=row,
                    row_number=current_row,
                    headers=headers,
                    column_mapping=column_mapping,
                )

                result.rows.append(parsed_row)
                rows_parsed += 1

        finally:
            workbook.close()

    def _parse_row(
        self,
        row: Any,  # openpyxl row tuple
        row_number: int,
        headers: list[str],
        column_mapping: dict[str, str] | None,
    ) -> ParsedRow:
        """Parse a single Excel row into a ParsedRow."""
        errors: list[FileParseError] = []
        data: dict[str, Any] = {}
        type_info: dict[str, str] = {}

        for i, cell in enumerate(row):
            if i < len(headers):
                col_name = headers[i]
            else:
                col_name = f"column_{i}"

            # Get cell value and type
            value = cell.value
            original_type = type(value).__name__ if value is not None else "None"
            type_info[col_name] = original_type

            # Convert to string for consistency
            string_value = self._cell_to_string(cell)
            data[col_name] = string_value

        # Apply column mapping
        data = self._apply_column_mapping(data, column_mapping)

        return ParsedRow(
            row_number=row_number,
            data=data,
            metadata={"original_types": type_info},
            errors=errors,
        )

    def _cell_to_string(self, cell: Any) -> str | None:
        """Convert Excel cell value to string."""
        if cell.value is None:
            return None

        value = cell.value

        # Handle different types
        if isinstance(value, datetime):
            return value.isoformat()
        elif isinstance(value, bool):
            return str(value).lower()
        elif isinstance(value, (int, float, Decimal)):
            # Preserve numeric precision
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            return str(value)
        else:
            return str(value).strip()
